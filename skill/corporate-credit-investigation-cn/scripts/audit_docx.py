import argparse
from collections import Counter
import json
import re
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


PLACEHOLDER_PATTERNS = {
    "待补充标记": r"【\s*(?:待补充|待核验|填写|金额|期限|业务品种|受信客户全称)[^】]*】",
    "英文占位符": r"(?i)(?<![A-Za-z])x{2,}(?![A-Za-z])",
    "粘贴说明": r"粘贴处|粘贴在这|待定",
}
VAGUE_PHRASES = ("以实际为准", "以最终为准", "有望", "具备一定", "需持续关注", "待核验")
TEMPLATE_RISK_PATTERNS = {
    "手机号": r"(?<!\d)1[3-9]\d{9}(?!\d)",
    "身份证号": r"(?<!\d)\d{17}[0-9Xx](?!\d)",
    "统一社会信用代码": r"(?<![0-9A-Z])[0-9A-HJ-NPQRTUWXY]{18}(?![0-9A-Z])",
    "固定年度": r"(?<!\d)(?:19|20)\d{2}年",
    "已勾选复选框": r"[☑☒■]",
    "预填判断": r"无异常|经营、财务状况[^。；]*良好|确认贸易背景真实|最新主体评级\s*[：:]\s*[A-D][0-9]|拟同意给予[^。；\n]{0,120}低风险授信额度[^。；\n]{0,120}授信方案有效期",
    "填写说明": r"这里要|需要按照[^。；\n]{0,80}(?:添加|粘贴)",
}
REPORT_SENSITIVE_PATTERNS = {
    key: TEMPLATE_RISK_PATTERNS[key]
    for key in ("手机号", "身份证号", "统一社会信用代码")
}
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
STORY_PART = re.compile(r"^word/(?:document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$")


def xml_story_text(path):
    """读取正文、页眉页脚、脚注和批注，避免遗漏旧客户信息。"""
    chunks = []
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not STORY_PART.match(name):
                continue
            try:
                root = ElementTree.fromstring(archive.read(name))
            except ElementTree.ParseError:
                continue
            chunks.append(f"[{name}]\n" + "".join(node.text or "" for node in root.iter(f"{{{WORD_NAMESPACE}}}t")))
    return "\n".join(chunks)


def regex_hits(patterns, text):
    return {
        label: len(re.findall(pattern, text, flags=re.IGNORECASE))
        for label, pattern in patterns.items()
        if re.search(pattern, text, flags=re.IGNORECASE)
    }


def package_stats(path):
    with ZipFile(path) as archive:
        names = set(archive.namelist())

        def count_parts(prefix):
            return len([name for name in names if name.startswith(prefix) and not name.endswith("/")])

        comments = 0
        comment_ids = []
        if "word/comments.xml" in names:
            root = ElementTree.fromstring(archive.read("word/comments.xml"))
            nodes = root.findall(f".//{{{WORD_NAMESPACE}}}comment")
            comments = len(nodes)
            comment_ids = [node.get(f"{{{WORD_NAMESPACE}}}id") for node in nodes]
        anchors = []
        ends = []
        if "word/document.xml" in names:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
            anchors = [node.get(f"{{{WORD_NAMESPACE}}}id") for node in root.findall(f".//{{{WORD_NAMESPACE}}}commentRangeStart")]
            ends = [node.get(f"{{{WORD_NAMESPACE}}}id") for node in root.findall(f".//{{{WORD_NAMESPACE}}}commentRangeEnd")]
        return {
            "comments": comments,
            "comment_anchors": len(anchors),
            "comment_ids": comment_ids,
            "comment_anchor_ids": anchors,
            "comment_end_ids": ends,
            "headers": count_parts("word/header"),
            "footers": count_parts("word/footer"),
            "media_files": count_parts("word/media/"),
            "embedded_objects": count_parts("word/embeddings/"),
        }


def numbering_stats(path):
    """列出自动编号段落及其编号定义，避免把不可见编号误当成普通文本。"""
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names or "word/numbering.xml" not in names:
            return []
        document_root = ElementTree.fromstring(archive.read("word/document.xml"))
        numbering_root = ElementTree.fromstring(archive.read("word/numbering.xml"))

    abstract_by_num = {}
    for node in numbering_root.findall(f".//{{{WORD_NAMESPACE}}}num"):
        num_id = node.get(f"{{{WORD_NAMESPACE}}}numId")
        abstract = node.find(f"{{{WORD_NAMESPACE}}}abstractNumId")
        if num_id is not None and abstract is not None:
            abstract_by_num[num_id] = abstract.get(f"{{{WORD_NAMESPACE}}}val")

    level_defs = {}
    for abstract in numbering_root.findall(f".//{{{WORD_NAMESPACE}}}abstractNum"):
        abstract_id = abstract.get(f"{{{WORD_NAMESPACE}}}abstractNumId")
        for level in abstract.findall(f"{{{WORD_NAMESPACE}}}lvl"):
            ilvl = level.get(f"{{{WORD_NAMESPACE}}}ilvl", "0")
            values = {}
            for tag in ("start", "numFmt", "lvlText"):
                child = level.find(f"{{{WORD_NAMESPACE}}}{tag}")
                values[tag] = child.get(f"{{{WORD_NAMESPACE}}}val") if child is not None else None
            level_defs[(abstract_id, ilvl)] = values

    numbered = []
    for index, paragraph in enumerate(document_root.findall(f".//{{{WORD_NAMESPACE}}}p")):
        num_pr = paragraph.find(f"./{{{WORD_NAMESPACE}}}pPr/{{{WORD_NAMESPACE}}}numPr")
        if num_pr is None:
            continue
        num_id_node = num_pr.find(f"{{{WORD_NAMESPACE}}}numId")
        ilvl_node = num_pr.find(f"{{{WORD_NAMESPACE}}}ilvl")
        if num_id_node is None:
            continue
        num_id = num_id_node.get(f"{{{WORD_NAMESPACE}}}val")
        ilvl = ilvl_node.get(f"{{{WORD_NAMESPACE}}}val") if ilvl_node is not None else "0"
        abstract_id = abstract_by_num.get(num_id)
        numbered.append({
            "paragraph": index,
            "text": paragraph_text(paragraph).strip(),
            "num_id": num_id,
            "level": ilvl,
            "abstract_num_id": abstract_id,
            **level_defs.get((abstract_id, ilvl), {}),
        })
    return numbered


def paragraph_text(element):
    return "".join(node.text or "" for node in element.iter(f"{{{WORD_NAMESPACE}}}t"))


def template_semantic_issues(document):
    issues = []
    tables = document.tables
    checks = {
        (12, 0, 0): "保证人间接融资情况",
        (12, 2, 0): "保证人直接融资情况",
        (21, 8, 0): "受信客户融资评价",
        (21, 9, 0): "受信客户融资变动原因分析",
    }
    for (table_index, row_index, cell_index), wrong_text in checks.items():
        if table_index >= len(tables):
            continue
        table = tables[table_index]
        if row_index >= len(table.rows) or cell_index >= len(table.rows[row_index].cells):
            continue
        if wrong_text in table.rows[row_index].cells[cell_index].text:
            issues.append(f"表{table_index + 1}第{row_index + 1}行存在对象错置标题“{wrong_text}”")
    if any(paragraph.text.strip() == "。" for paragraph in all_paragraphs(document)):
        issues.append("模板存在孤立句号段落")
    return issues


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def run_format_stats(document):
    fonts, sizes, bold = Counter(), Counter(), Counter()
    text_runs = 0
    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            text_runs += 1
            props = run._r.rPr
            mappings = [props.rFonts.get(qn(f"w:{key}")) or "继承" for key in ("ascii", "hAnsi", "eastAsia", "cs")] if props is not None and props.rFonts is not None else ["继承"] * 4
            fonts["|".join(mappings)] += 1
            sizes[str(run.font.size.pt if run.font.size else "继承")] += 1
            bold[str(run.bold if run.bold is not None else "继承")] += 1
    return {"text_runs": text_runs, "font_mappings": dict(fonts.most_common()), "font_sizes_pt": dict(sizes.most_common()), "bold_values": dict(bold.most_common())}


def dominant_key(values):
    return max(values, key=values.get) if values else None


def document_stats(document):
    paragraphs = [p.text for p in document.paragraphs]
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    return {"paragraphs": len(document.paragraphs), "tables": len(document.tables), "images": len(document.inline_shapes), "sections": len(document.sections), "characters": len("\n".join(paragraphs + cells))}


def main():
    parser = argparse.ArgumentParser(description="检查授信报告DOCX的客户残留、占位符、预填结论和分页结构。")
    parser.add_argument("document")
    parser.add_argument("--mode", choices=("report", "template"), default="report")
    parser.add_argument("--strict", action="store_true", help="发现问题或结构警告时返回非零状态。")
    parser.add_argument("--baseline", help="用户修改基准版，用于检测内容或对象数量异常减少。")
    parser.add_argument("--forbidden", action="append", default=[])
    parser.add_argument("--forbidden-file", action="append", default=[], help="每行一个历史客户名、机构名或其他禁用文本；空行和#注释忽略。")
    parser.add_argument("--output")
    args = parser.parse_args()

    path = Path(args.document)
    document = Document(path)
    paragraphs = [p.text for p in document.paragraphs]
    full_text = xml_story_text(path)
    normalized_text = re.sub(r"\s+", "", full_text)
    forbidden_words = list(args.forbidden)
    for forbidden_file in args.forbidden_file:
        forbidden_path = Path(forbidden_file)
        if not forbidden_path.is_file():
            parser.error(f"禁用词文件不存在：{forbidden_path}")
        forbidden_words.extend(
            line.strip() for line in forbidden_path.read_text(encoding="utf-8").splitlines()
            if line.strip() and not line.lstrip().startswith("#")
        )
    leftovers = {}
    for word in dict.fromkeys(forbidden_words):
        normalized_word = re.sub(r"\s+", "", word)
        if normalized_word and normalized_word in normalized_text:
            leftovers[word] = normalized_text.count(normalized_word)
    placeholders = regex_hits(PLACEHOLDER_PATTERNS, full_text)
    # 封面的“（2018年版）”等版本标识属于模板元数据，不是客户报告期。
    # 只豁免完整括号内的“四位年份+年版”，正文固定年度仍按原规则拦截。
    template_risk_text = re.sub(r"[（(](?:19|20)\d{2}年版[）)]", "", full_text)
    template_risks = regex_hits(TEMPLATE_RISK_PATTERNS, template_risk_text) if args.mode == "template" else {}
    sensitive_data = regex_hits(REPORT_SENSITIVE_PATTERNS, full_text) if args.mode == "report" else {}
    vague = {word: full_text.count(word) for word in VAGUE_PHRASES if word in full_text}
    empty_headings = [{"paragraph": i, "text": value.strip()} for i, value in enumerate(paragraphs) if re.fullmatch(r"[（(]?[一二三四五六七八九十0-9]+[）).、]", value.strip())]

    page_break_before, keep_next, keep_lines = [], [], []
    explicit_breaks = 0
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.paragraph_format.page_break_before:
            page_break_before.append(index)
        if paragraph.paragraph_format.keep_with_next:
            keep_next.append(index)
        if paragraph.paragraph_format.keep_together:
            keep_lines.append(index)
        explicit_breaks += len(paragraph._p.xpath('.//w:br[@w:type="page"]'))

    cant_split_rows = fixed_height_rows = 0
    for table in document.tables:
        for row in table.rows:
            props = row._tr.find(qn("w:trPr"))
            if props is not None and props.find(qn("w:cantSplit")) is not None:
                cant_split_rows += 1
            if props is not None and props.find(qn("w:trHeight")) is not None:
                fixed_height_rows += 1

    issues = []
    if leftovers:
        issues.append(f"存在{sum(leftovers.values())}处禁用文本")
    if args.mode == "report" and placeholders:
        issues.append(f"正式报告存在{sum(placeholders.values())}处占位或填写说明")
    if args.mode == "template" and template_risks:
        issues.append(f"模板存在{sum(template_risks.values())}处敏感信息、固定时点或预填判断")
    if args.mode == "template":
        issues.extend(template_semantic_issues(document))

    report = {
        "document": str(path), "mode": args.mode, "stats": document_stats(document),
        "package_stats": package_stats(path), "run_format_stats": run_format_stats(document),
        "explicit_page_breaks": explicit_breaks, "page_break_before_paragraphs": page_break_before,
        "keep_with_next_paragraphs": keep_next, "keep_together_paragraphs": keep_lines,
        "cant_split_table_rows": cant_split_rows, "fixed_height_table_rows": fixed_height_rows,
        "forbidden_leftovers": leftovers, "placeholders": placeholders, "template_risks": template_risks,
        "sensitive_data_hits": sensitive_data, "automatic_numbering": numbering_stats(path),
        "vague_phrases": vague, "empty_heading_like": empty_headings, "issues": issues, "warnings": [],
        "note": "结构检查不能替代逐页渲染。automatic_numbering记录Word自动编号定义，普通文本提取不含渲染后的编号。模板模式允许明确占位符和封面“四位年份+年版”版本标识，但不允许真实个人信息、固定客户报告期、已勾选选项或预填结论。",
    }

    if args.baseline:
        baseline_path = Path(args.baseline)
        baseline = Document(baseline_path)
        baseline_stats = document_stats(baseline)
        baseline_package = package_stats(baseline_path)
        baseline_formats = run_format_stats(baseline)
        report.update({"baseline": str(baseline_path), "baseline_stats": baseline_stats, "baseline_package_stats": baseline_package, "baseline_run_format_stats": baseline_formats})
        for key in ("paragraphs", "tables", "images", "sections", "characters"):
            if report["stats"][key] < baseline_stats[key]:
                report["warnings"].append(f"{key}由{baseline_stats[key]}减少为{report['stats'][key]}，请确认未误删内容")
        for key in ("comments", "comment_anchors", "headers", "footers", "media_files", "embedded_objects"):
            if report["package_stats"][key] < baseline_package[key]:
                report["warnings"].append(f"{key}由{baseline_package[key]}减少为{report['package_stats'][key]}，请确认对象未丢失")
        for key in ("comment_ids", "comment_anchor_ids", "comment_end_ids"):
            missing = [item for item in baseline_package[key] if item not in report["package_stats"][key]]
            if missing:
                report["warnings"].append(f"{key}缺少基准ID: {', '.join(missing)}")
            elif report["package_stats"][key] != baseline_package[key]:
                report["warnings"].append(f"{key}顺序发生变化，请逐条核对批注锚点")
        for key, label in (("font_mappings", "主字体映射"), ("font_sizes_pt", "主字号"), ("bold_values", "主要字重")):
            if dominant_key(baseline_formats[key]) != dominant_key(report["run_format_stats"][key]):
                report["warnings"].append(f"{label}发生变化，请确认未发生全局格式漂移")

    output = Path(args.output) if args.output else path.with_suffix(".文档核验.json")
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "issues": issues, "warnings": report["warnings"]}, ensure_ascii=False))
    if args.strict and (issues or report["warnings"]):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
