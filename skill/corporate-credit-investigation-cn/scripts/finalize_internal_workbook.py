"""补充artifact-tool尚未覆盖的Excel打印与工作表保护设置。"""

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Protection
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    args = parser.parse_args()
    path = Path(args.workbook)
    workbook = load_workbook(path)
    sheet = workbook.active

    sheet.print_area = "A1:H23"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_view.selection[0].activeCell = "A1"
    sheet.sheet_view.selection[0].sqref = "A1"

    for row in sheet.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)
    for row in sheet["C4:E11"]:
        for cell in row:
            cell.protection = Protection(locked=False)
    for row in sheet["E13:E19"]:
        for cell in row:
            cell.protection = Protection(locked=False)

    if not any("E17" in rule.sqref for rule in sheet.data_validations.dataValidation):
        rule = DataValidation(type="decimal", operator="between", formula1=1, formula2=2, allow_blank=True)
        rule.promptTitle = "调节系数"
        rule.prompt = "请输入1至2之间的数值，一般取1。"
        rule.errorTitle = "输入超出范围"
        rule.error = "调节系数必须在1至2之间。"
        rule.showInputMessage = True
        rule.showErrorMessage = True
        sheet.add_data_validation(rule)
        rule.add(sheet["E17"])

    sheet.protection.sheet = True
    sheet.protection.selectLockedCells = True
    sheet.protection.selectUnlockedCells = False
    if workbook.calculation is None:
        workbook.calculation = CalcProperties(calcMode="auto")
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)


if __name__ == "__main__":
    main()
