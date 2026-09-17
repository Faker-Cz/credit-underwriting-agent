import argparse
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook


NUMBER_PATTERN = re.compile(r"^[+-]?(?:\d{1,3}(?:[,，]\d{3})+|\d+)(?:\.\d+)?%?$")
MIXED_AMOUNT_RATIO_PATTERN = re.compile(r"^[+-]?(?:\d{1,3}(?:[,，]\d{3})+|\d+)(?:\.\d+)?\s*[（(][^）)]*[）)]$")
BALANCE_PAIRS = (("资产总计", "负债和所有者权益总计"), ("资产总计", "负债及所有者权益总计"))
LOAN_FORMULAS = {
    "F4": "=(C4+D4)/2", "G4": '=IF(F4=0,"",$D$9/F4)', "H4": '=IF(OR(G4="",G4=0),"",360/G4)',
    "F5": "=(C5+D5)/2", "G5": '=IF(F5=0,"",$D$10/F5)', "H5": '=IF(OR(G5="",G5=0),"",360/G5)',
    "F6": "=(C6+D6)/2", "G6": '=IF(F6=0,"",$D$10/F6)', "H6": '=IF(OR(G6="",G6=0),"",360/G6)',
    "F7": "=(C7+D7)/2", "G7": '=IF(F7=0,"",$D$10/F7)', "H7": '=IF(OR(G7="",G7=0),"",360/G7)',
    "F8": "=(C8+D8)/2", "G8": '=IF(F8=0,"",$D$9/F8)', "H8": '=IF(OR(G8="",G8=0),"",360/G8)',
    "F9": "=(C9+D9)/2", "F10": "=(C10+D10)/2", "F11": "=(C11+D11)/2",
    "D12": '=IF(C9=0,"",(D9-C9)/C9)',
    "E20": '=IF((SUM(H4:H6)-SUM(H7:H8))=0,"",360/(SUM(H4:H6)-SUM(H7:H8)))',
    "E21": '=IF(OR(D9="",D9=0,E20="",E20=0),"",D9*(1-(D9-D10)/D9)*(1+E13)/E20)',
    "E22": '=IF(OR(E21="",E17=""),"",E21*E17+E18+E19-E14-E15-E16)',
}


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def normalized_label(value):
    return re.sub(r"\s+", "", str(value or "")).replace("：", "").replace(":", "")


def find_labels(sheet):
    labels, duplicates = {}, []
    for row in sheet.iter_rows():
        for cell in row[: min(3, len(row))]:
            label = normalized_label(cell.value)
            if not label or label in {"—", "-"}:
                continue
            if label in labels and labels[label] != cell.row:
                duplicates.append({"label": label, "rows": [labels[label], cell.row]})
            labels.setdefault(label, cell.row)
    return labels, duplicates


def numeric_columns(sheet, first_row, second_row):
    checks = []
    for column in range(2, sheet.max_column + 1):
        first, second = sheet.cell(first_row, column).value, sheet.cell(second_row, column).value
        if is_number(first) and is_number(second):
            difference = first - second
            checks.append({"column": column, "first": first, "second": second, "difference": difference, "status": abs(difference) <= 0.01})
    return checks


def validation_covers(sheet, coordinate, minimum=None, maximum=None):
    for rule in sheet.data_validations.dataValidation:
        if coordinate not in rule.sqref:
            continue
        if minimum is not None and str(rule.formula1) != str(minimum):
            continue
        if maximum is not None and str(rule.formula2) != str(maximum):
            continue
        return True
    return False


def audit_loan_template(sheet, report):
    for coordinate, expected in LOAN_FORMULAS.items():
        actual = sheet[coordinate].value
        if actual != expected:
            report["issues"].append(f"{sheet.title}!{coordinate}: 关键公式不符，期望 {expected}，实际 {actual}")
        if not sheet[coordinate].protection.locked:
            report["issues"].append(f"{sheet.title}!{coordinate}: 公式单元格未锁定")
    for row in range(4, 12):
        for column in range(3, 6):
            if sheet.cell(row, column).protection.locked:
                report["issues"].append(f"{sheet.title}!{sheet.cell(row, column).coordinate}: 黄色输入区应解锁")
    for row in range(13, 20):
        if sheet.cell(row, 5).protection.locked:
            report["issues"].append(f"{sheet.title}!E{row}: 黄色输入区应解锁")
    if not sheet.protection.sheet:
        report["issues"].append(f"{sheet.title}: 工作表保护未启用，锁定公式不能生效")
    if "A1:H23" not in str(sheet.print_area).replace("$", ""):
        report["issues"].append(f"{sheet.title}: 打印区域应覆盖 A1:H23")
    if not sheet.sheet_properties.pageSetUpPr.fitToPage or sheet.page_setup.fitToWidth != 1 or sheet.page_setup.fitToHeight != 1:
        report["issues"].append(f"{sheet.title}: 未设置为单页宽高打印")
    if not validation_covers(sheet, "E17", 1, 2):
        report["issues"].append(f"{sheet.title}!E17: 缺少1至2之间的数据验证")
    for typo in ("其他融资融资", "分别是是"):
        if any(typo in str(cell.value or "") for row in sheet.iter_rows() for cell in row):
            report["issues"].append(f"{sheet.title}: 存在文字错误“{typo}”")


def main():
    parser = argparse.ArgumentParser(description="审计授信财务汇总及流贷测算工作簿。")
    parser.add_argument("workbook")
    parser.add_argument("--mode", choices=("generic", "loan-template"), default="generic")
    parser.add_argument("--strict", action="store_true", help="发现问题时返回非零状态。")
    parser.add_argument("--output")
    parser.add_argument("--pure-values", action="store_true", help="将任何公式视为交付错误。")
    args = parser.parse_args()

    path = Path(args.workbook)
    workbook = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    report = {"workbook": str(path), "mode": args.mode, "pure_values": args.pure_values, "sheets": {}, "issues": [], "external_links": len(getattr(workbook, "_external_links", []))}
    if report["external_links"]:
        report["issues"].append(f"工作簿存在{report['external_links']}个外部链接")

    for sheet in workbook.worksheets:
        labels, duplicate_labels = find_labels(sheet)
        text_numbers, mixed_cells, formulas, errors = [], [], [], []
        cached_sheet = cached[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    formulas.append(cell.coordinate)
                    cached_value = cached_sheet[cell.coordinate].value
                    if isinstance(cached_value, str) and cached_value.startswith("#"):
                        errors.append({"cell": cell.coordinate, "value": cached_value})
                elif isinstance(value, str):
                    stripped = value.strip()
                    if NUMBER_PATTERN.fullmatch(stripped):
                        text_numbers.append({"cell": cell.coordinate, "value": stripped})
                    if MIXED_AMOUNT_RATIO_PATTERN.fullmatch(stripped):
                        mixed_cells.append({"cell": cell.coordinate, "value": stripped})
                if cell.data_type == "e":
                    errors.append({"cell": cell.coordinate, "value": value})
        hidden_rows = [index for index, dimension in sheet.row_dimensions.items() if dimension.hidden]
        hidden_columns = [key for key, dimension in sheet.column_dimensions.items() if dimension.hidden]
        sheet_report = {
            "rows": sheet.max_row, "columns": sheet.max_column,
            "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
            "duplicate_labels": duplicate_labels, "numeric_values_stored_as_text": text_numbers,
            "mixed_amount_ratio_cells": mixed_cells, "formula_cells": formulas, "error_cells": errors,
            "hidden_rows": hidden_rows, "hidden_columns": hidden_columns, "sheet_state": sheet.sheet_state,
        }
        for first_label, second_label in BALANCE_PAIRS:
            first, second = normalized_label(first_label), normalized_label(second_label)
            if first in labels and second in labels:
                checks = numeric_columns(sheet, labels[first], labels[second])
                sheet_report["balance_checks"] = checks
                if any(not item["status"] for item in checks):
                    report["issues"].append(f"{sheet.title}: 资产负债表不平")
                break
        if duplicate_labels:
            report["issues"].append(f"{sheet.title}: 存在{len(duplicate_labels)}组重复标签")
        if text_numbers:
            report["issues"].append(f"{sheet.title}: 存在{len(text_numbers)}个疑似文本数字")
        if mixed_cells:
            report["issues"].append(f"{sheet.title}: 存在{len(mixed_cells)}个金额与比例/横线混写单元格")
        if errors:
            report["issues"].append(f"{sheet.title}: 存在{len(errors)}个Excel错误值")
        if args.pure_values and formulas:
            report["issues"].append(f"{sheet.title}: 纯数值交付仍含{len(formulas)}个公式")
        if sheet.sheet_state != "visible":
            report["issues"].append(f"{sheet.title}: 存在隐藏工作表")
        report["sheets"][sheet.title] = sheet_report

    if args.mode == "loan-template":
        if len(workbook.worksheets) != 1:
            report["issues"].append("流贷清洁模板应仅含一个工作表")
        audit_loan_template(workbook.active, report)

    output = Path(args.output) if args.output else path.with_suffix(".财务核验.json")
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "issues": report["issues"]}, ensure_ascii=False))
    if args.strict and report["issues"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
