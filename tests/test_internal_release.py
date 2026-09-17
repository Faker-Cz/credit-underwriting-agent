import hashlib
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "skill" / "corporate-credit-investigation-cn"
ASSETS = SKILL / "assets"


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class InternalReleaseTests(unittest.TestCase):
    def test_root_and_skill_assets_match(self):
        for name in ("单一客户授信调查报告模版.docx", "低风险报告模版.docx", "流动资金类贷款授信额度测算表.xlsx"):
            self.assertEqual(sha256(ROOT / name), sha256(ASSETS / name), name)

    def test_docx_templates_are_clean(self):
        script = SKILL / "scripts" / "audit_docx.py"
        with tempfile.TemporaryDirectory() as directory:
            for name in ("单一客户授信调查报告模版.docx", "低风险报告模版.docx"):
                result = subprocess.run(
                    [sys.executable, str(script), str(ASSETS / name), "--mode", "template", "--strict", "--forbidden", "浙商银行", "--forbidden", "龙华支行", "--output", str(Path(directory) / f"{name}.json")],
                    capture_output=True,
                    text=True,
                )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_loan_workbook_contract(self):
        script = SKILL / "scripts" / "audit_financial_workbook.py"
        with tempfile.TemporaryDirectory() as directory:
            result = subprocess.run(
                [sys.executable, str(script), str(ASSETS / "流动资金类贷款授信额度测算表.xlsx"), "--mode", "loan-template", "--strict", "--output", str(Path(directory) / "audit.json")],
                capture_output=True,
                text=True,
            )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

        workbook = load_workbook(ASSETS / "流动资金类贷款授信额度测算表.xlsx", data_only=False)
        sheet = workbook.active
        self.assertFalse(sheet["E12"].protection.locked)
        self.assertFalse(sheet.protection.selectLockedCells)

    def test_template_semantic_corrections(self):
        standard = Document(ASSETS / "单一客户授信调查报告模版.docx")
        checks = {
            (12, 0, 0): "受信客户间接融资情况",
            (12, 2, 0): "受信客户直接融资情况",
            (21, 8, 0): "保证人融资评价",
            (21, 9, 0): "保证人融资变动原因分析",
        }
        for (table_index, row_index, cell_index), expected in checks.items():
            self.assertIn(expected, standard.tables[table_index].rows[row_index].cells[cell_index].text)

        low_risk = Document(ASSETS / "低风险报告模版.docx")
        text = "\n".join(paragraph.text for paragraph in low_risk.paragraphs)
        text += "\n" + "\n".join(cell.text for table in low_risk.tables for row in table.rows for cell in row.cells)
        self.assertNotIn("授信方案有效期1年，业务品种为流动性支持类", text)
        self.assertFalse(any(paragraph.text.strip() == "。" for paragraph in low_risk.paragraphs))

    def test_report_forbidden_file_is_enforced(self):
        script = SKILL / "scripts" / "audit_docx.py"
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            report_path = directory / "report.docx"
            forbidden_path = directory / "forbidden.txt"
            output_path = directory / "audit.json"
            document = Document()
            document.add_paragraph("历史客户甲")
            document.save(report_path)
            forbidden_path.write_text("# 项目残留词\n历史客户甲\n", encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(script), str(report_path), "--mode", "report", "--strict", "--forbidden-file", str(forbidden_path), "--output", str(output_path)],
                capture_output=True,
                text=True,
            )
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("存在1处禁用文本", result.stdout)

    def test_maintenance_rules_do_not_embed_project_period_or_counts(self):
        references = SKILL / "references"
        combined = "\n".join(path.read_text(encoding="utf-8") for path in references.glob("*.md"))
        self.assertNotIn("2023—最新一期", combined)
        self.assertNotIn("审计合并范围22个主体", combined)
        self.assertNotIn("直接投资16家、控制企业23家", combined)

    def test_canonical_skill_name(self):
        guide = (ROOT / "授信报告智能体使用说明与操作指引.md").read_text(encoding="utf-8")
        self.assertNotIn("bank-corporate-credit-investigation-cn", guide)
        self.assertIn("$corporate-credit-investigation-cn", guide)
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        self.assertIn("--baseline", readme)
        self.assertIn("--pure-values", readme)


if __name__ == "__main__":
    unittest.main()
