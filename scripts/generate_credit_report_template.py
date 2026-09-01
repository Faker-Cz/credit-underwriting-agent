import argparse
from pathlib import Path
from itertools import count

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT_CN = "SimSun"
FONT_HEAD = "SimHei"
PLACEHOLDER = RGBColor(128, 128, 128)
BLACK = RGBColor(0, 0, 0)
CHECKBOX_IDS = count(100001)


def set_font(run, name=FONT_CN, size=12, bold=False, color=BLACK):
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().rFonts
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(slot), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_checkbox_control(paragraph):
    """Append a clickable Word checkbox content control (unchecked -> checked)."""
    control_id = next(CHECKBOX_IDS)
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")

    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), f"授信选项{control_id}")
    sdt_pr.append(alias)

    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), f"credit_checkbox_{control_id}")
    sdt_pr.append(tag)

    sdt_id = OxmlElement("w:id")
    sdt_id.set(qn("w:val"), str(control_id))
    sdt_pr.append(sdt_id)

    checkbox = OxmlElement("w14:checkbox")
    checked = OxmlElement("w14:checked")
    checked.set(qn("w14:val"), "0")
    checkbox.append(checked)

    checked_state = OxmlElement("w14:checkedState")
    checked_state.set(qn("w14:val"), "2611")
    checked_state.set(qn("w14:font"), "MS Gothic")
    checkbox.append(checked_state)

    unchecked_state = OxmlElement("w14:uncheckedState")
    unchecked_state.set(qn("w14:val"), "2610")
    unchecked_state.set(qn("w14:font"), "MS Gothic")
    checkbox.append(unchecked_state)
    sdt_pr.append(checkbox)
    sdt.append(sdt_pr)

    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        run_fonts.set(qn(f"w:{slot}"), "MS Gothic")
    run_pr.append(run_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    run_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "24")
    run_pr.append(size_cs)
    run.append(run_pr)
    text = OxmlElement("w:t")
    text.text = "☐"
    run.append(text)
    sdt_content.append(run)
    sdt.append(sdt_content)
    paragraph._p.append(sdt)


def add_mixed_text(paragraph, text, *, name=FONT_CN, size=12, bold=False, color=BLACK):
    """Write normal text and replace each □ marker with a clickable checkbox."""
    parts = text.split("□")
    for index, part in enumerate(parts):
        if part:
            run = paragraph.add_run(part)
            set_font(run, name=name, size=size, bold=bold, color=color)
        if index < len(parts) - 1:
            add_checkbox_control(paragraph)


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, size=8):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "000000")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def cell_text(cell, text="", *, bold=False, center=False, size=12, placeholder=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    # 用户明确要求表格内全部采用宋体小四；保留 size 参数仅为兼容既有调用。
    add_mixed_text(p, text, size=12, bold=bold, color=PLACEHOLDER if placeholder else BLACK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    return cell


def merge_text(table, r1, c1, r2, c2, text, *, bold=False, center=True, size=12, placeholder=False):
    c = table.cell(r1, c1).merge(table.cell(r2, c2))
    cell_text(c, text, bold=bold, center=center, size=size, placeholder=placeholder)
    return c


def make_table(doc, rows, cols, widths=None, font_size=12):
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                if i < len(row.cells):
                    set_cell_width(row.cells[i], width)
    for row in table.rows:
        prevent_row_split(row)
        for cell in row.cells:
            cell_text(cell, "", size=font_size)
    return table


def para(doc, text="", *, size=12, bold=False, center=False, indent=0, before=0, after=0, keep=False, font=FONT_CN, placeholder=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    # 不写入“与下段同页”等分页属性，避免在显示格式标记时出现标题前黑点。
    # 正文最低字号统一为宋体小四（12 磅）；标题和封面较大字号保持原设定。
    add_mixed_text(p, text, name=font, size=max(12, size), bold=bold, color=PLACEHOLDER if placeholder else BLACK)
    return p


def heading(doc, text, level=1):
    sizes = {1: 22, 2: 16, 3: 16, 4: 16}
    p = para(doc, text, size=sizes.get(level, 16), bold=False, before=6 if level == 1 else 3, after=3, keep=False, font=FONT_HEAD)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def box(doc, label, prompt="{{填写}}", height_lines=3, size=12):
    t = make_table(doc, 1, 1, [17.2], font_size=size)
    c = t.cell(0, 0)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_font(r, size=12, bold=True)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(max(10, height_lines * 10))
    r2 = p2.add_run(prompt)
    set_font(r2, size=12, color=PLACEHOLDER)
    return t


def add_footer_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("—")
    set_font(r, size=9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r2 = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r2.append(t)
    fld.append(r2)
    p._p.append(fld)
    r3 = p.add_run("—")
    set_font(r3, size=9)


def set_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.65)
    add_footer_page_number(section)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    rfonts = normal._element.get_or_add_rPr().rFonts
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(slot), FONT_CN)
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15


def add_cover(doc):
    para(doc, "附件1", size=12)
    para(doc, "", after=28)
    para(doc, "浙  商  银  行", size=20, center=True, font=FONT_CN)
    para(doc, "对公客户授信调查报告", size=22, center=True, font=FONT_CN)
    para(doc, "", after=42)
    labels = ["受信客户：", "申报单位：", "联系人：", "联系电话：", "填报时间："]
    for label in labels:
        p = para(doc, "", size=14, indent=4.0, after=2)
        r1 = p.add_run(label)
        set_font(r1, size=14)
        r2 = p.add_run("____________________")
        set_font(r2, size=14)


def add_scheme_and_evaluation(doc):
    heading(doc, "第一部分 授信申请与评价", 1)
    heading(doc, "一、授信方案申请（参照决策意见单描述）", 2)
    para(doc, "参考格式：", size=12, bold=True)
    para(doc, "申请给予{{客户名称}}授信类型：□一般授信  □特别授信  □专项授信；额度{{金额及币种}}，额度有效期{{期限}}，额度控制方式：□可循环  □不可循环；业务品种{{品种}}，担保方式{{担保方式}}。", size=12, placeholder=True)
    para(doc, "放款条件：{{决策意见单原文}}", size=12, placeholder=True)
    para(doc, "管理要求：{{决策意见单原文}}", size=12, placeholder=True)
    heading(doc, "二、授信方案评价", 2)
    heading(doc, "（一）合规性评价", 3)
    t = make_table(doc, 7, 1, [17.2])
    cell_text(t.cell(0, 0), "本次申请授信用途：\n□维持正常经营需要  □生产规模扩大或结算方式变更  □置换他行授信\n□季节性生产需求  □特定项目建设或投资  □其他", bold=True, size=9)
    cell_text(t.cell(1, 0), "信用风险限额与资金需求测算：\n1.信用风险限额测算结果：{{填写}}\n2.新增流动资金需求测算结果：{{填写；未申请流动性支持类产品可不填}}", bold=True, size=9, placeholder=False)
    merge_text(t, 2, 0, 3, 0, "资金用途、还款来源与额度合理性说明（如申请额度超过信用风险限额测算结果60%的，须说明理由）：\n{{填写}}", bold=True, center=False, size=9, placeholder=False)
    merge_text(t, 4, 0, 6, 0, "其他合规问题与异常事项（如行业准入、经营资格、涉案涉诉、环保税务违法、生产安全等）：\n{{填写}}", bold=True, center=False, size=9, placeholder=False)
    heading(doc, "（二）申请人基本面评价", 3)
    b = make_table(doc, 18, 5, [4.5, 3.15, 3.15, 3.15, 3.25])
    merge_text(b, 0, 0, 0, 4, "经营概况", bold=True, size=9.5)
    cell_text(b.cell(1, 0), "所属行业\n（国标大-中-小类）", bold=True, center=True, size=9)
    merge_text(b, 1, 1, 1, 4, "{{国标大类—中类—小类}}", placeholder=True, size=9.5)
    cell_text(b.cell(2, 0), "主营业务及行业地位", bold=True, center=True, size=9)
    merge_text(b, 2, 1, 2, 4, "{{填写}}", placeholder=True, size=9.5)
    merge_text(b, 3, 0, 3, 4, "主要财务数据  单位：{{万元/亿元}}", bold=True, size=9.5)
    headers = ["科目", "{{前三年}}", "{{前两年}}", "{{前一年}}", "{{最近一期}}"]
    for j, v in enumerate(headers):
        cell_text(b.cell(4, j), v, bold=True, center=True, size=8.8, placeholder=v.startswith("{{"))
    metrics = ["总资产", "总负债", "刚性负债", "营业收入", "净利润", "经营净现金流", "净现金流", "总资产报酬率", "资产负债率", "流动比率"]
    for i, m in enumerate(metrics, start=5):
        cell_text(b.cell(i, 0), m, bold=True, center=True, size=8.8)
        for j in range(1, 5):
            cell_text(b.cell(i, j), "{{填}}", center=True, size=8.5, placeholder=True)
    merge_text(b, 15, 0, 16, 4, "重要风险指标：\n1.是否参与跨行业、跨产业链投资：□否；□是，涉及行业{{填写}}\n2.上年末刚性负债结构：金融机构负债占比{{填}}；非银金融机构负债占比{{填}}；其他占比{{填}}。\n3.上年度经营净现金流/财务费用={{填}}\n风险指标异常说明：{{填写}}", center=False, size=8.8)
    merge_text(b, 17, 0, 17, 4, "营销背景及合作情况（前次授信批复及投放情况、与本机构合作时间、历次授信额度变化、其他业务开展情况、投放效益等）：\n{{填写}}", bold=True, center=False, size=8.8)
    heading(doc, "（三）授信方案总体评价", 3)
    box(doc, "（包括但不限于授信理由、授信方案的合理性、可操作性、风险管控的有效性及综合效益）", "{{填写}}", 4)


def add_declaration(doc):
    heading(doc, "三、调查声明与承诺", 2)
    para(doc, "郑重承诺：", size=12)
    para(doc, "（一）主协办调查人采用本《调查报告》及本机构相关调查指引所要求的调查方式、调查内容与调查要求对受信人、保证人、抵质押物及具体项目等进行了调查，调查过程如下：", size=12, indent=0.6)
    para(doc, "调查日期：{{填写}}\n本机构参加人员：{{填写}}\n调查访问对象及职务：{{填写}}", size=12, indent=1.1, placeholder=True)
    para(doc, "（二）主协办调查人对本调查报告所陈述事实、数据和授信资料的真实性、完整性、有效性承担责任，并愿意按照本机构相关管理办法承担相应的调查责任。", size=12, indent=0.6)
    para(doc, "（三）部门负责人已对主协办客户经理的调查行为履行充分的监督责任。", size=12, indent=0.6)
    para(doc, "主办调查人（签字）：________\n协办调查人（签字）：________\n部门负责人（签字）：________", size=12, indent=10.2)


def add_basic_info(doc, prefix="受信客户"):
    heading(doc, "第二部分 调查内容", 1)
    heading(doc, "一、受信客户调查分析", 2)
    heading(doc, "（一）基本情况", 3)
    t = make_table(doc, 16, 4, [3.25, 5.0, 3.25, 5.0])
    merge_text(t, 0, 0, 0, 3, f"{prefix}基本情况", bold=True, size=9.5)
    pairs = [
        ("客户全称", "所处行业"), ("注册资本", "注册地址"), ("实收资本", "法定代表人"),
        ("经营范围", "实际控制人"), ("企业性质", "营业执照号码"), ("基本开户行", "机构代码证号码"),
        ("信用等级", "贷款卡号或中征码")
    ]
    for i, (a, b) in enumerate(pairs, start=1):
        cell_text(t.cell(i, 0), a, center=True, bold=True, size=8.8)
        cell_text(t.cell(i, 1), "{{填写}}", placeholder=True, size=8.8)
        cell_text(t.cell(i, 2), b, center=True, bold=True, size=8.8)
        cell_text(t.cell(i, 3), "{{填写}}", placeholder=True, size=8.8)
    headers = ["主要股东", "投资金额", "股权占比", "实际出资"]
    for j, h in enumerate(headers):
        cell_text(t.cell(8, j), h, center=True, bold=True, size=8.8)
    for r in range(9, 12):
        for c in range(4):
            cell_text(t.cell(r, c), "{{填写}}" if r == 9 else "", placeholder=True, size=8.5)
    merge_text(t, 12, 0, 12, 3, "1.受信客户历史沿革（如发展过程、股权变化等）：\n{{填写}}", center=False, size=9)
    merge_text(t, 13, 0, 13, 3, "2.主要股东、实际控制人（法定代表人）情况介绍：\n{{填写}}", center=False, size=9)
    merge_text(t, 14, 0, 14, 3, "所处集团基本情况", bold=True, size=9.5)
    merge_text(t, 15, 0, 15, 3, "1.所处集团、集团本级及核心子公司简介（附组织架构图）\n2.集团主营收入结构及业务板块\n3.集团、集团本级及核心子公司基本财务情况\n{{填写}}", center=False, size=9)


def add_nonfinancial(doc):
    heading(doc, "（二）受信客户非财务因素分析", 3)
    t = make_table(doc, 1, 1, [17.2])
    text = (
        "受信客户经营情况\n"
        "1.经营团队情况（重要成员介绍、团队经营管理能力及稳定性评价等）\n{{填写}}\n\n"
        "2.经营策略及发展战略\n{{填写}}\n\n"
        "3.客户所处行业（发展现状及趋势、竞争格局等）及行业地位（市场份额及排名等）\n{{填写}}\n\n"
        "4.经营要素情况（经营场地、员工构成、经营方式、生产工艺、技术及研发、产能规模及产能利用率等）\n{{填写}}\n\n"
        "5.主营业务（或产品）的收入结构及毛利率\n{{按前3年、前2年、前1年、当期填写收入、占比及毛利率}}\n\n"
        "6.成本及采购\n（1）产品成本构成、原材料采购模式、结算方式及账期等\n{{填写}}\n"
        "（2）前五大供应商情况（关联请说明）：供应商、上年度采购额、当期采购额、当期占比、结算方式、账期\n{{填写}}\n\n"
        "7.销售\n（1）销售模式、产品价格趋势、结算方式及账期等\n{{填写}}\n"
        "（2）前五大销售商情况（关联请说明）：销售商、上年度销售额、当期销售额、当期占比、结算方式、账期\n{{填写}}"
    )
    cell_text(t.cell(0, 0), text, size=9)
    para(doc, "", after=2)
    t2 = make_table(doc, 1, 1, [17.2])
    cell_text(t2.cell(0, 0), "受信客户对外投资情况\n1.受信客户重大投资项目情况（投资规模、资金安排、项目进度、效益测算等）\n（1）受信客户在建、拟建项目情况：{{填写}}\n（2）受信客户重大对外投资情况（股权、房地产等）：{{填写}}\n2.实际控制人及其控制的关联企业对外投资情况：{{填写}}", size=9)


def add_financial(doc, subject="受信客户", guarantee=False):
    title = "3.保证人财务因素分析" if guarantee else "（三）受信客户财务因素分析"
    heading(doc, title, 3 if not guarantee else 2)
    heading(doc, "（1）报表合规性与真实性说明" if guarantee else "1.报表合规性与真实性", 4)
    a = make_table(doc, 5, 4, [4.2, 4.3, 4.3, 4.4])
    merge_text(a, 0, 0, 0, 3, "财务报表是否审计  □是  □否", size=9)
    headers = ["项目", "{{前3年}}", "{{前2年}}", "{{前1年}}"]
    for j, h in enumerate(headers):
        cell_text(a.cell(1, j), h, center=True, bold=True, size=8.8, placeholder=h.startswith("{{"))
    for i, label in enumerate(["审计单位", "非标意见及备注"], start=2):
        cell_text(a.cell(i, 0), label, center=True, bold=True, size=8.8)
        for j in range(1, 4):
            cell_text(a.cell(i, j), "{{填写}}", placeholder=True, size=8.5)
    merge_text(a, 4, 0, 4, 3, "若未提供审计报告，说明原因：{{填写}}", center=False, size=8.8)
    heading(doc, "（2）保证人合并报表范围（如有）" if guarantee else "2.合并报表范围", 4)
    box(doc, "合并报表范围", "{{逐一列示纳入合并范围的主体名称、持股比例、取得方式及变动情况}}", 3)
    heading(doc, "（3）保证人财务状况和财务指标变动表（粘贴处）" if guarantee else "3.企业财务状况和财务指标变动表（粘贴处）", 4)
    box(doc, "财务报表及指标表粘贴区域", "{{插入资产负债表、利润表、现金流量表及财务指标变动表}}", 6)
    # 小四字号下该标题容易孤置于页底，采用显式分页，不使用标题分页属性。
    page_break(doc)
    heading(doc, "（4）保证人财务分析" if guarantee else "4.财务分析", 4)
    analysis = (
        "（1）重要科目明细及重大变化分析（包括但不限于：应收账款、预付账款、其他应收款、存货、长期股权投资、固定资产、无形资产、应付账款、预收账款、其他应付款、资本公积，以及金额占比超过15%、余额变化超过20%的重要科目）\n{{填写}}\n\n"
        "（2）主营收入和主营业务利润率，三项费用（销售、管理、财务费用），投资收益、营业外收入等非经常性损益分析\n{{填写}}\n\n"
        "（3）现金流量分析\n{{填写}}\n\n"
        "（4）财务分析（偿债能力、营运能力、盈利能力、成长能力等）\n{{填写}}\n\n"
        "（5）其他需要说明事项（如资产质量、收入真实性的核查和评价、税收优惠等）\n{{填写}}"
    )
    box(doc, "", analysis, 8, size=9)


def add_core_assets(doc):
    heading(doc, "（四）受信客户核心资产梳理（主要指房地产、股权等）", 3)
    t = make_table(doc, 12, 7, [2.7, 2.7, 2.2, 2.1, 1.7, 2.4, 3.4])
    merge_text(t, 0, 0, 0, 6, "单位：{{万元}}", size=9.5)
    sections = [(1, "受信客户报表内核心资产"), (5, "受信客户报表外核心资产"), (8, "受信客户实际控制人拥有资产（含关联企业）")]
    for row, title in sections:
        merge_text(t, row, 0, row, 6, title, bold=True, size=9)
        headers = ["资产品种", "权证号/编码", "数量/面积", "账面价值", "估值", "抵质押情况", "使用情况"]
        for j, h in enumerate(headers):
            cell_text(t.cell(row + 1, j), h, center=True, bold=True, size=8)
    cell_text(t.cell(3, 0), "房地产（坐落）", center=True, size=8)
    cell_text(t.cell(4, 0), "股权/其他", center=True, size=8)
    for r in (3, 4, 7, 10):
        for c in range(1, 7):
            cell_text(t.cell(r, c), "{{填}}" if r in (3, 7, 10) else "", placeholder=True, size=8)
    merge_text(t, 11, 0, 11, 6, "补充说明：{{填写}}", center=False, size=8.8)


def financing_block(doc, subject="受信客户"):
    heading(doc, f"（五）{subject}融资及或有负债分析" if subject == "受信客户" else f"4.{subject}融资及或有负债分析", 3 if subject == "受信客户" else 2)
    t = make_table(doc, 19, 7, [2.7, 3.1, 2.15, 1.25, 3.1, 2.2, 2.7])
    merge_text(t, 0, 0, 0, 6, f"{subject}间接融资情况", bold=True, size=9.5)
    merge_text(t, 1, 0, 1, 6, f"（1）{subject}间接融资情况表    数据时间：{{年/月末}}；单位：{{万元}}", center=False, size=8.8)
    hdr = ["类别", "授信机构名称", "余额", "敞口", "担保方式/担保单位", "期限/到期日", "备注"]
    for j, h in enumerate(hdr):
        cell_text(t.cell(2, j), h, center=True, bold=True, size=7.8)
    labels = ["短期借款", "应付票据", "长期借款", "长期应付款等其他科目", "敞口合计"]
    for i, lab in enumerate(labels, start=3):
        cell_text(t.cell(i, 0), lab, center=True, size=8)
        for j in range(1, 7):
            cell_text(t.cell(i, j), "{{填}}" if i < 7 else "", placeholder=True, size=8)
    merge_text(t, 8, 0, 8, 6, f"（2）{subject}间接融资变动表    数据时间：{{年/月末}}；单位：{{万元}}", center=False, size=8.8)
    hdr2 = ["授信机构名称", "当期授信额度", "当期敞口", "上年末敞口", "前年末敞口", "担保方式", "备注"]
    for j, h in enumerate(hdr2):
        cell_text(t.cell(9, j), h, center=True, bold=True, size=7.8)
    for i, label in enumerate(["金融机构1", "金融机构2", "非银金融机构1", "非银金融机构2", "其他", "合计"], start=10):
        cell_text(t.cell(i, 0), label, center=True, size=8)
    merge_text(t, 16, 0, 16, 6, f"{subject}直接融资情况    数据时间：{{年/月末}}；单位：{{万元}}", bold=True, size=8.8)
    merge_text(t, 17, 0, 17, 6, "融资品种、代码、融资余额、利率、发行日、到期日、增信方式：{{填写}}", center=False, size=8.8)
    merge_text(t, 18, 0, 18, 6, f"{subject}最新主体评级情况（评级机构、评级时间、评级结果、评级展望及近三期迁徙变化）：{{填写}}", center=False, size=8.8)
    q = make_table(doc, 8, 1, [17.2])
    cell_text(q.cell(0, 0), "中征码查询信息", bold=True, center=True, size=9.2)
    cell_text(q.cell(1, 0), "1.人行查询日期与结果分析，当期信贷余额与查询结果对比差异简析：{{填写}}\n2.查询异常结果说明（逾期、五级分类为关注及以下、曾经有过不良信用记录等）：{{填写}}", size=8.8)
    cell_text(q.cell(2, 0), f"{subject}融资评价", bold=True, center=True, size=9.2)
    cell_text(q.cell(3, 0), f"1.{subject}融资变动原因分析：{{填写}}\n2.{subject}融资结构、融资稳定性、融资能力评价，关注大额融资到期情况：{{填写}}", size=8.8)
    cell_text(q.cell(4, 0), "或有负债分析", bold=True, center=True, size=9.2)
    cell_text(q.cell(5, 0), "1.内部关联担保情况：{{填写}}\n2.外部担保情况及大额被担保单位风险评判：{{填写}}", size=8.8)
    cell_text(q.cell(6, 0), "外部担保明细：被担保方、担保形式、担保金额、到期日、关系、代偿可能性、是否互保及金额。", size=8.8)
    cell_text(q.cell(7, 0), "3.当期担保情况与中征码查询结果对比差异分析：{{填写}}", size=8.8)


def add_advantages_and_guarantee(doc):
    heading(doc, "（六）受信客户优劣势分析", 3)
    box(doc, "（从财务因素、非财务因素等方面对受信客户的优势、劣势进行分析）", "{{填写}}", 4)
    heading(doc, "二、专业调查模块（勾选，同时填写特定报告模板作为附件）", 2)
    para(doc, "□股票质押式回购  □可交换（可转换）债券  □配资类\n□永续债权  □资产证券化  □股权投资  □项目融资\n□并购贷款  □债券承销  □其他", size=11)
    heading(doc, "三、担保（增信）分析", 2)
    heading(doc, "（一）抵（质）押担保", 3)
    box(doc, "1.抵（质）押物调查", "{{抵（质）押物基本信息、价值评估、抵质押状态、使用状态及市场情况}}", 4)
    box(doc, "2.抵（质）押人担保意愿及担保能力分析", "{{抵质押人简介、与受信人关系、可处置变现能力等}}", 3)


def add_guarantor(doc):
    heading(doc, "（二）保证人（增信方）调查分析：{{保证人名称}}", 3)
    heading(doc, "1.保证人基本情况", 2)
    t = make_table(doc, 15, 4, [3.25, 5.0, 3.25, 5.0])
    merge_text(t, 0, 0, 0, 3, "保证人基本情况", bold=True, size=9.5)
    pairs = [("客户全称", "所处行业"), ("注册资本", "注册地址"), ("实收资本", "法定代表人"), ("经营范围", "实际控制人"), ("企业性质", "营业执照号码"), ("基本开户行", "机构代码证号码"), ("信用等级", "贷款卡号或中征码")]
    for i, (a, b) in enumerate(pairs, start=1):
        cell_text(t.cell(i, 0), a, center=True, bold=True, size=8.7)
        cell_text(t.cell(i, 1), "{{填写}}", placeholder=True, size=8.5)
        cell_text(t.cell(i, 2), b, center=True, bold=True, size=8.7)
        cell_text(t.cell(i, 3), "{{填写}}", placeholder=True, size=8.5)
    for j, h in enumerate(["主要股东", "投资金额", "股权占比", "实际出资"]):
        cell_text(t.cell(8, j), h, center=True, bold=True, size=8.5)
    for r in range(9, 12):
        for c in range(4):
            cell_text(t.cell(r, c), "{{填写}}" if r == 9 else "", placeholder=True, size=8.3)
    merge_text(t, 12, 0, 12, 3, "（1）保证人历史沿革（如发展过程、股权变化等）：{{填写}}", center=False, size=8.8)
    merge_text(t, 13, 0, 13, 3, "（2）主要股东、实际控制人（法定代表人）情况介绍（如为集团内保，此处可简略）：{{填写}}", center=False, size=8.8)
    merge_text(t, 14, 0, 14, 3, "（3）所处集团基本情况、经营收入结构、集团合并、本级及核心子公司基本财务信息（如为集团内保，此处可简略）：{{填写}}", center=False, size=8.8)
    heading(doc, "2.保证人非财务因素分析", 2)
    box(doc, "（1）保证人经营要素、产购销情况", "{{填写}}", 3)
    box(doc, "（2）核心资产及对外投资情况", "{{填写}}", 3)


def add_guarantor_tail(doc):
    financing_block(doc, "保证人")
    heading(doc, "5.保证人优劣势分析", 2)
    box(doc, "（从财务因素、非财务因素等方面对保证人的优势、劣势进行分析，含担保意愿及担保能力）", "{{填写}}", 4)
    heading(doc, "四、授信方案风险管控措施", 2)
    box(doc, "（包括授信方案设计、业务/产品安排、放款条件设定、业务存续期检查与监测要求及其他授信后管理措施等）", "{{原则上按决策意见单原文填写；凡方案字段发生调整，应同步更新额度、期限、品种、循环方式、担保安排、放款条件、管理要求及所有联动表述}}", 7)


def build(output_path):
    doc = Document()
    configure_styles(doc)
    set_page(doc)
    add_cover(doc)

    page_break(doc)
    add_scheme_and_evaluation(doc)

    page_break(doc)
    add_declaration(doc)

    page_break(doc)
    add_basic_info(doc)

    page_break(doc)
    add_nonfinancial(doc)

    page_break(doc)
    add_financial(doc)

    page_break(doc)
    add_core_assets(doc)

    page_break(doc)
    financing_block(doc)

    page_break(doc)
    add_advantages_and_guarantee(doc)

    page_break(doc)
    add_guarantor(doc)

    page_break(doc)
    add_financial(doc, subject="保证人", guarantee=True)

    page_break(doc)
    add_guarantor_tail(doc)

    core = doc.core_properties
    core.title = "对公客户授信调查报告模板"
    core.subject = "授信调查报告可复用模板"
    core.author = ""
    core.keywords = "授信调查报告, 对公客户, 模板"

    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="生成通用对公客户授信调查报告模板。")
    parser.add_argument("--output", required=True, help="由使用者自行指定的本地DOCX输出路径。")
    build(parser.parse_args().output)
